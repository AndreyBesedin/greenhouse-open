"""The 2024 build, from a miniature timeseries archive: one compartment's
readings, the site's weather and forecast, and the harvest workbook become
one canonical greenhouse, the same bytes every time."""

import zipfile
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from greenhouse_protocol.enums import EventType, ObservationType

from greenhouse_adapters.canonical import CanonicalGreenhouse
from greenhouse_adapters.manifests import load_manifest
from greenhouse_adapters.storage.layout import DataDirectory
from greenhouse_adapters.storage.resolver import ArtifactResolver
from greenhouse_adapters.wur.agc4_challenge_2024 import (
    DATASET,
    TIMESERIES_ARTIFACT,
    TIMESERIES_MEMBER_PREFIX,
)
from greenhouse_adapters.wur.agc4_challenge_2024.build import HARVEST_MEMBER, build_greenhouse
from greenhouse_adapters.wur.agc4_challenge_2024.channels import FORECAST_MEMBER, WEATHER_MEMBER
from greenhouse_adapters.wur.agc4_challenge_2024.compartments import COMPARTMENTS, compartment

REFERENCE = compartment("3.06")
# The plants are spaced out on 1 October, and harvested on day 320, 15
# November, at noon.
REFERENCE_CSV = """\
time,compartment/air_temperature,dwarf_tomato/plant_density,dwarf_tomato/harvest_date
2024-09-03 00:00:00+02:00,20.0,40.0,
2024-10-01 00:00:00+02:00,21.0,20.0,
2024-11-15 12:00:00+01:00,19.0,20.0,320.0
"""
# The site's first reading comes before the compartment's.
WEATHER_CSV = """\
time,weather/air_temperature.outside
2024-09-02 23:55:00+02:00,15.0
"""
FORECAST_CSV = """\
time,weather_forecast/radiation_sum
2024-09-03 00:00:00+02:00,1200.0
"""


def _harvest_workbook() -> bytes:
    book = openpyxl.Workbook()
    info = book.active
    assert info is not None
    info.title = "Information"
    sampling = book.create_sheet("Destructive harvest 7-10-24")
    sampling.append(["plant", "FW ", "SE FW", "# tom", "SE #tom"])
    sampling.append(["306-1", 60.0, None, 30, None])
    final = book.create_sheet("Final Harvest  3.06")
    final.append(
        ["GH #", "plant", "FW tot", "# tot", "# tom green", "fw green", "#tom red", "FW red"]
    )
    final.append([306, "306-1", 300.0, 40, 10, 50.0, 30, 250.0])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def data_dir(tmp_path: Path) -> DataDirectory:
    manifest = load_manifest(DATASET.id)
    artifact = manifest.artifact(TIMESERIES_ARTIFACT)
    data_dir = DataDirectory(tmp_path)
    path = data_dir.raw(manifest) / artifact.name
    path.parent.mkdir(parents=True)
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(TIMESERIES_MEMBER_PREFIX + REFERENCE.timeseries_member, REFERENCE_CSV)
        archive.writestr(TIMESERIES_MEMBER_PREFIX + WEATHER_MEMBER, WEATHER_CSV)
        archive.writestr(TIMESERIES_MEMBER_PREFIX + FORECAST_MEMBER, FORECAST_CSV)
        archive.writestr(TIMESERIES_MEMBER_PREFIX + HARVEST_MEMBER, _harvest_workbook())
    # the fixture is not the real artifact: pre-mark it verified, as the
    # resolver records a verification, with what names the file as it is
    status = path.stat()
    path.with_name(path.name + ".md5-verified").write_text(
        f"{artifact.md5} {status.st_size} {status.st_ino} {status.st_mtime_ns}\n"
    )
    return data_dir


def _build(data_dir: DataDirectory) -> CanonicalGreenhouse:
    return build_greenhouse(data_dir, ArtifactResolver(data_dir), [REFERENCE])


def test_a_build_is_the_same_bytes_every_time(data_dir: DataDirectory) -> None:
    first = _build(data_dir).provenance().content_sha256
    second = _build(data_dir).provenance().content_sha256

    assert second == first


def test_the_greenhouse_lists_every_compartment_and_is_dated_by_the_first_reading(
    data_dir: DataDirectory,
) -> None:
    built = _build(data_dir)
    greenhouse = built.greenhouse()

    assert [c.compartment_id for c in greenhouse.compartments] == [
        c.compartment_id for c in COMPARTMENTS.values()
    ]
    assert greenhouse.created_at == datetime(2024, 9, 2, 21, 55, tzinfo=UTC)
    provenance = built.provenance()
    assert provenance.selection == {"compartments": ["3.06"], "teams": [REFERENCE.team]}
    assert [s.member for s in provenance.sources] == [
        TIMESERIES_MEMBER_PREFIX + member
        for member in (REFERENCE.timeseries_member, WEATHER_MEMBER, FORECAST_MEMBER, HARVEST_MEMBER)
    ]


def test_the_compartments_readings_are_its_own_and_the_sites_are_the_greenhouses(
    data_dir: DataDirectory,
) -> None:
    observations = list(_build(data_dir).observations())
    scope = {o.observation_type: o.compartment_id for o in observations}

    assert scope[ObservationType.AIR_TEMPERATURE_C] == REFERENCE.compartment_id
    assert scope[ObservationType.SAMPLED_FRUIT_COUNT_PER_PLANT] == REFERENCE.compartment_id
    assert scope[ObservationType.OUTSIDE_AIR_TEMPERATURE_C] is None
    assert scope[ObservationType.FORECAST_RADIATION_SUM_TODAY_J_CM2] is None


def test_the_spacing_and_the_final_harvest_are_events(data_dir: DataDirectory) -> None:
    events = {e.event_type: e for e in _build(data_dir).events()}

    assert set(events) == {EventType.SPACING, EventType.HARVEST}
    assert events[EventType.SPACING].timestamp == datetime(2024, 9, 30, 22, tzinfo=UTC)
    # noon on 15 November, in Amsterdam's winter time
    assert events[EventType.HARVEST].timestamp == datetime(2024, 11, 15, 11, tzinfo=UTC)
