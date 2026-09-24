import zipfile
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

import openpyxl
import pytest
from greenhouse_protocol.enums import ObservationType, SourceType

from greenhouse_adapters.manifests import load_manifest
from greenhouse_adapters.storage.layout import DataDirectory
from greenhouse_adapters.storage.resolver import ArtifactResolver
from greenhouse_adapters.wur.agc4_pretrial_2023 import (
    CLIMATE_MEMBER,
    CROP_MEMBER,
    DESTRUCTIVE_MEMBER,
    TIMESERIES_ARTIFACT,
)
from greenhouse_adapters.wur.agc4_pretrial_2023.build import build_greenhouse
from greenhouse_adapters.wur.agc4_pretrial_2023.climate import parse_climate
from greenhouse_adapters.wur.agc4_pretrial_2023.compartments import COMPARTMENT_ID, GREENHOUSE_ID
from greenhouse_adapters.wur.datasets import AGC4_PRETRIAL_2023

# The real sheet's columns, including the ones that must not be ingested.
HEADER = [
    "datenum", "date", "time", "tout", "rhout", "iglob", "windsp", "rain", "parout",
    "pyrgeo", "co2out", "t_air", "rh", "co2", "scr_enrg", "scr_blck", "ligth_on",
    "vent_lee", "vent_wind", "t_rail", "part1", "par2", "par3", "par4",
]  # fmt: skip
SEP_5 = 739134  # datenum of 5 September 2023
OCT_29 = SEP_5 + 54  # clocks go back: 02:00-02:59 local happens twice


def _row(datenum: float, **values: object) -> list[object]:
    cells: dict[str, object] = {column: "NaN" for column in HEADER}
    cells.update(datenum=datenum, date=None, time=None)
    cells.update(values)
    return [cells[column] for column in HEADER]


def _workbook(*rows: list[object], header: list[str] = HEADER) -> bytes:
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "weather_climate"
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    book.create_sheet("Info")
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


ROWS = (
    _row(SEP_5),  # the export's leading all-NaN padding
    _row(
        SEP_5 + 0.5,
        tout=18.5,
        co2out=412.0,
        rain=0.4,
        t_air=21.0,
        co2=650.0,
        t_rail=38.5,
        part1=210.0,
        ligth_on=1.0,
    ),
    _row(OCT_29 + 2.5 / 24),  # the repeated DST hour, left empty by the export
    _row(OCT_29 + 3.5 / 24, t_air=17.0),
)


def test_site_weather_and_compartment_climate_are_scoped_apart() -> None:
    observations = list(parse_climate(BytesIO(_workbook(*ROWS))))

    noon = {
        o.observation_type: o
        for o in observations
        if o.timestamp == datetime(2023, 9, 5, 10, tzinfo=UTC)
    }
    # zone PAR (part1) and the undocumented lamp column are not ingested
    assert set(noon) == {
        ObservationType.OUTSIDE_AIR_TEMPERATURE_C,
        ObservationType.OUTSIDE_CO2_PPM,
        ObservationType.OUTSIDE_RAIN,
        ObservationType.AIR_TEMPERATURE_C,
        ObservationType.CO2_PPM,
        ObservationType.HEATING_PIPE_TEMPERATURE_C,
    }
    assert noon[ObservationType.OUTSIDE_AIR_TEMPERATURE_C].compartment_id is None
    assert noon[ObservationType.OUTSIDE_CO2_PPM].value == 412.0
    assert noon[ObservationType.OUTSIDE_RAIN].value == 0.4
    assert noon[ObservationType.AIR_TEMPERATURE_C].compartment_id == COMPARTMENT_ID
    assert noon[ObservationType.HEATING_PIPE_TEMPERATURE_C].value == 38.5
    assert all(o.greenhouse_id == GREENHOUSE_ID and o.plant_id is None for o in observations)
    assert all(o.source.type == SourceType.IMPORTED_DATA for o in observations)
    assert len({o.observation_id for o in observations}) == len(observations)


def test_empty_rows_are_skipped_even_inside_the_repeated_dst_hour() -> None:
    observations = list(parse_climate(BytesIO(_workbook(*ROWS))))

    assert sorted({o.timestamp for o in observations}) == [
        datetime(2023, 9, 5, 10, tzinfo=UTC),
        # 03:30 CET, after the clocks went back
        datetime(2023, 10, 29, 2, 30, tzinfo=UTC),
    ]


def test_a_reading_inside_the_repeated_dst_hour_is_refused() -> None:
    ambiguous = _row(OCT_29 + 2.5 / 24, t_air=17.2)

    with pytest.raises(ValueError, match="occurs twice"):
        list(parse_climate(BytesIO(_workbook(ambiguous))))


def test_a_sheet_missing_an_expected_column_is_refused() -> None:
    header = [column for column in HEADER if column != "t_rail"]

    with pytest.raises(ValueError, match="t_rail"):
        list(parse_climate(BytesIO(_workbook(header=header))))


def test_build_writes_one_greenhouse_with_its_single_compartment(tmp_path: Path) -> None:
    manifest = load_manifest(AGC4_PRETRIAL_2023.id)
    artifact = manifest.artifact(TIMESERIES_ARTIFACT)
    data_dir = DataDirectory(tmp_path)
    path = data_dir.raw(manifest) / artifact.name
    path.parent.mkdir(parents=True)
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(CLIMATE_MEMBER, _workbook(*ROWS))
        archive.writestr(CROP_MEMBER, _empty_crop_workbook())
        archive.writestr(DESTRUCTIVE_MEMBER, _empty_destructive_workbook())
    # the fixture is not the real artifact: pre-mark it verified
    path.with_name(path.name + ".md5-verified").write_text(artifact.md5 + "\n")

    first = build_greenhouse(data_dir, ArtifactResolver(data_dir))
    second = build_greenhouse(data_dir, ArtifactResolver(data_dir))

    assert second.provenance().content_sha256 == first.provenance().content_sha256
    greenhouse = first.greenhouse()
    assert greenhouse.greenhouse_id == GREENHOUSE_ID
    assert [c.compartment_id for c in greenhouse.compartments] == [COMPARTMENT_ID]
    assert greenhouse.created_at == datetime(2023, 9, 5, 10, tzinfo=UTC)
    assert first.provenance().observation_count == 7
    assert [s.member for s in first.provenance().sources] == [
        CLIMATE_MEMBER,
        CROP_MEMBER,
        DESTRUCTIVE_MEMBER,
    ]
    assert first.provenance().event_count == 0
    assert greenhouse.compartments[0].plants == []


def _empty_crop_workbook() -> bytes:
    """A crop-measurement workbook with the expected columns and no plants."""
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "All data"
    identity = ["Date", "Week", "# on label", "Variety", "Corrected EC", "Light", "field"]
    single = ["Plant repetition#", "Plantheigth", "# leaves", "leaf length", "leaf width"]
    trusses = [
        f"# {kind} trus{t}"
        for t in range(1, 6)
        for kind in ("flowering", "set", "yellow", "orange", "red")
    ]
    sheet.append(identity + single + trusses + ["#trusses"])
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _empty_destructive_workbook() -> bytes:
    """A destructive-harvest workbook with the expected header and no samples."""
    book = openpyxl.Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "All Data"
    sheet.append(["Sowing date", None, "2023-08-15"])
    sheet.append(
        ["DAS", "Date", "Phase", "Plant density (p/m2)", "Treatment", "Variety", "EC", "Light"]
        + ["Sample name", "Sample n", "Plant height"]
    )
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()
