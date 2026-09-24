from datetime import UTC, datetime
from io import BytesIO

import openpyxl
from greenhouse_protocol.enums import EventSource, EventType, ObservationType

from greenhouse_adapters.wur.agc4_challenge_2024.compartments import compartment
from greenhouse_adapters.wur.agc4_challenge_2024.harvest import (
    final_harvest_event,
    read_harvest_workbook,
    sampling_observations,
)


def _workbook() -> BytesIO:
    """A miniature Harvest.xlsx with the real sheets' shapes: a first
    sampling with the short header, a later one with the long header and
    a compartment aggregate row, and a final-harvest sheet."""
    book = openpyxl.Workbook()
    info = book.active
    assert info is not None
    info.title = "Information"
    info.append(["Legend", "Description", "unit"])

    first = book.create_sheet("Destructive harvest 7-10-24")
    first.append(["plant", "FW ", "SE FW", "# tom", "SE #tom"])
    first.append(["306-1", 60.0, None, 30, None])
    first.append(["306-2", 80.0, None, 40, None])
    first.append(["301-1", 10.0, None, 1, None])
    first.append([306, 70.0, 10.0, 35, 5.0])  # aggregate row: ignored

    later = book.create_sheet("Destructive harvest 22-10-24")
    later.append(["GH #", "plant", "FW tot", "#tot", "# tom green", "FW green", "#tom red"])
    later.append([306, "306-1", 200.0, 50, 50, 200.0, 0])
    later.append([306, None, 200.0, 50, 50, 200.0, 0])  # aggregate row: ignored

    final = book.create_sheet("Final Harvest  3.06")
    final.append(
        ["GH #", "plant", "FW tot", "# tot", "# tom green", "fw green", "#tom red", "FW red"]
    )
    final.append([306, "306-1", 300.0, 40, 10, 50.0, 30, 250.0])
    final.append([306, "306-2", 100.0, 20, 5, 20.0, 15, 80.0])

    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    return buffer


def test_samplings_become_mean_per_plant_observations_at_local_noon() -> None:
    workbook = read_harvest_workbook(_workbook())

    observations = list(sampling_observations(workbook, compartment("3.06")))

    by_key = {(o.timestamp, o.observation_type): o.value for o in observations}
    noon_oct_7 = datetime(2024, 10, 7, 10, tzinfo=UTC)
    assert by_key[(noon_oct_7, ObservationType.SAMPLED_FRUIT_COUNT_PER_PLANT)] == 35.0
    assert by_key[(noon_oct_7, ObservationType.SAMPLED_FRUIT_FRESH_WEIGHT_G_PER_PLANT)] == 70.0
    noon_oct_22 = datetime(2024, 10, 22, 10, tzinfo=UTC)
    assert by_key[(noon_oct_22, ObservationType.SAMPLED_FRUIT_COUNT_PER_PLANT)] == 50.0
    assert all(o.plant_id is None for o in observations)
    assert all(o.greenhouse_id == "wur_agc4_2024" for o in observations)
    assert all(o.compartment_id == "3.06" for o in observations)


def test_other_compartments_samples_never_leak_in() -> None:
    workbook = read_harvest_workbook(_workbook())

    observations = list(sampling_observations(workbook, compartment("3.01")))

    assert [o.value for o in observations] == [1.0, 10.0]
    assert list(sampling_observations(workbook, compartment("3.08"))) == []


def test_final_harvest_is_a_greenhouse_level_event_with_declared_date_source() -> None:
    workbook = read_harvest_workbook(_workbook())
    recording_end = datetime(2024, 11, 15, 11, tzinfo=UTC)

    event = final_harvest_event(workbook, compartment("3.06"), harvested_at=recording_end)

    assert event is not None
    assert event.event_type == EventType.HARVEST
    assert event.source == EventSource.HUMAN_REPORTED
    assert event.plant_id is None
    assert event.compartment_id == "3.06"
    assert event.timestamp == recording_end
    assert event.parameters["harvested_mass_g"] == 400.0
    assert event.parameters["harvested_fruit_count"] == 60
    assert event.parameters["plants_sampled"] == 2
    assert event.parameters["mean_red_fruit_count_per_plant"] == 22.5
    assert event.parameters["date_source"] == "end_of_recording"
    assert final_harvest_event(workbook, compartment("3.01"), harvested_at=recording_end) is None


def test_final_harvest_event_records_where_its_date_came_from() -> None:
    workbook = read_harvest_workbook(_workbook())
    harvested_at = datetime(2024, 11, 15, 11, tzinfo=UTC)

    event = final_harvest_event(
        workbook,
        compartment("3.06"),
        harvested_at=harvested_at,
        date_source="dwarf_tomato/harvest_date",
    )

    assert event is not None
    assert event.timestamp == harvested_at
    assert event.parameters["date_source"] == "dwarf_tomato/harvest_date"
