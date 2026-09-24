"""Harvest.xlsx -> canonical records.

The workbook holds manual, destructive measurements of a handful of sample
plants per compartment: three dated intermediate samplings (all
compartments on one sheet each) and one final-harvest sheet per
compartment. The dataset never shared these with the teams, so they are
outcome ground truth, not something a controller could have seen.

- an intermediate sampling becomes two compartment-level Observations per
  compartment (mean fruit count and mean fruit fresh weight per sampled
  plant), stamped local noon on the sheet's date;
- the final harvest becomes one compartment-level HARVEST Event per
  compartment. The workbook gives no date for it; the compartment's time
  series records the harvest day, so the caller supplies that instant (or,
  if a file lacks it, the last recorded timestamp) and the event says which
  in `date_source`.
"""

import re
from collections.abc import Callable, Iterator, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from statistics import mean
from typing import IO

import openpyxl
from greenhouse_protocol.enums import EventSource, EventType, ObservationType
from greenhouse_protocol.event import Event
from greenhouse_protocol.observation import Observation

from greenhouse_adapters.wur.agc4_challenge_2024.compartments import (
    GREENHOUSE_ID,
    Compartment,
    compartment_by_code,
)
from greenhouse_adapters.wur.agc4_challenge_2024.timeseries import SOURCE, observation_id
from greenhouse_adapters.wur.common.time import local_noon

_SAMPLING_SHEET = re.compile(r"^Destructive harvest (\d{1,2})-(\d{1,2})-(\d{2})$")
_FINAL_SHEET = re.compile(r"^Final Harvest\s+(\d\.\d{2})\s*$")
_PLANT_LABEL = re.compile(r"^(\d{3})-(\d+)$")


@dataclass(frozen=True)
class SampledPlant:
    compartment_code: str
    label: str
    fresh_weight_g: float
    fruit_count: int
    red_fruit_count: int | None
    red_fresh_weight_g: float | None


@dataclass(frozen=True)
class Sampling:
    day: date
    plants: list[SampledPlant]

    def for_compartment(self, code: str) -> list[SampledPlant]:
        return [p for p in self.plants if p.compartment_code == code]


@dataclass(frozen=True)
class FinalHarvest:
    compartment_code: str
    plants: list[SampledPlant]


@dataclass(frozen=True)
class HarvestWorkbook:
    samplings: list[Sampling]
    final_harvests: list[FinalHarvest]


def read_harvest_workbook(source: Path | IO[bytes]) -> HarvestWorkbook:
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    samplings: list[Sampling] = []
    finals: list[FinalHarvest] = []
    for sheet_name in workbook.sheetnames:
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        if match := _SAMPLING_SHEET.match(sheet_name.strip()):
            day_, month, year = (int(g) for g in match.groups())
            samplings.append(Sampling(date(2000 + year, month, day_), _sampled_plants(rows)))
        elif match := _FINAL_SHEET.match(sheet_name):
            code = match.group(1).replace(".", "")
            finals.append(FinalHarvest(code, _sampled_plants(rows)))
    workbook.close()
    samplings.sort(key=lambda s: s.day)
    return HarvestWorkbook(samplings=samplings, final_harvests=finals)


def sampling_observations(
    workbook: HarvestWorkbook, compartment: Compartment
) -> Iterator[Observation]:
    for sampling in workbook.samplings:
        plants = sampling.for_compartment(compartment.code)
        if not plants:
            continue
        timestamp = local_noon(sampling.day)
        values = {
            ObservationType.SAMPLED_FRUIT_COUNT_PER_PLANT: mean(p.fruit_count for p in plants),
            ObservationType.SAMPLED_FRUIT_FRESH_WEIGHT_G_PER_PLANT: mean(
                p.fresh_weight_g for p in plants
            ),
        }
        for observation_type, value in values.items():
            yield Observation(
                observation_id=observation_id(compartment, timestamp, observation_type.value),
                greenhouse_id=GREENHOUSE_ID,
                compartment_id=compartment.compartment_id,
                plant_id=None,
                timestamp=timestamp,
                observation_type=observation_type,
                value=float(value),
                source=SOURCE,
            )


def final_harvest_event(
    workbook: HarvestWorkbook,
    compartment: Compartment,
    *,
    harvested_at: datetime,
    date_source: str = "end_of_recording",
) -> Event | None:
    final = next(
        (f for f in workbook.final_harvests if f.compartment_code == compartment.code), None
    )
    if final is None or not final.plants:
        return None
    plants = final.plants
    red_counts = [p.red_fruit_count for p in plants if p.red_fruit_count is not None]
    red_weights = [p.red_fresh_weight_g for p in plants if p.red_fresh_weight_g is not None]
    return Event(
        event_id=f"wur24_c{compartment.code}_final_harvest",
        greenhouse_id=GREENHOUSE_ID,
        compartment_id=compartment.compartment_id,
        plant_id=None,
        timestamp=harvested_at,
        event_type=EventType.HARVEST,
        source=EventSource.HUMAN_REPORTED,
        parameters={
            "harvested_mass_g": sum(p.fresh_weight_g for p in plants),
            "harvested_fruit_count": sum(p.fruit_count for p in plants),
            "plants_sampled": len(plants),
            "mean_fresh_weight_g_per_plant": mean(p.fresh_weight_g for p in plants),
            "mean_fruit_count_per_plant": mean(p.fruit_count for p in plants),
            "mean_red_fruit_count_per_plant": mean(red_counts) if red_counts else None,
            "mean_red_fresh_weight_g_per_plant": mean(red_weights) if red_weights else None,
            "date_source": date_source,
        },
    )


_Row = tuple[object, ...]


def _sampled_plants(rows: Sequence[_Row]) -> list[SampledPlant]:
    if not rows:
        return []
    header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
    plant_col = _find(header, ("plant",))
    weight_col = _find(header, ("fw tot", "fw"))
    count_col = _find(header, ("# tot", "#tot", "# tom"))
    red_count_col = _find(header, ("#tom red", "# tom red"), required=False)
    red_weight_col = _find(header, ("fw red", "fw tom red"), required=False)

    plants: list[SampledPlant] = []
    for row in rows[1:]:
        label = _cell(row, plant_col)
        match = _PLANT_LABEL.match(str(label).strip()) if label is not None else None
        if match is None:
            continue  # aggregate / blank rows
        code = match.group(1)
        if compartment_by_code(code) is None:
            continue
        weight = _number(_cell(row, weight_col))
        count = _number(_cell(row, count_col))
        if weight is None or count is None:
            continue
        red_count = _number(_cell(row, red_count_col)) if red_count_col is not None else None
        red_weight = _number(_cell(row, red_weight_col)) if red_weight_col is not None else None
        plants.append(
            SampledPlant(
                compartment_code=code,
                label=str(label).strip(),
                fresh_weight_g=weight,
                fruit_count=int(round(count)),
                red_fruit_count=int(round(red_count)) if red_count is not None else None,
                red_fresh_weight_g=red_weight,
            )
        )
    return plants


def _find(header: list[str], names: tuple[str, ...], *, required: bool = True) -> int | None:
    for name in names:
        if name in header:
            return header.index(name)
    if required:
        raise ValueError(f"harvest sheet has no column named any of {names}; header={header}")
    return None


def _cell(row: _Row, index: int | None) -> object:
    if index is None or index >= len(row):
        return None
    return row[index]


def _number(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return float(value)
    try:
        return float(str(value).strip())
    except ValueError:
        return None


ReadWorkbook = Callable[[Path], HarvestWorkbook]
